from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from django.test import SimpleTestCase
from docx import Document as WordDocument

from apps.knowledge.extractors.docx import DocxExtractionError, extract_docx

from .docx_fixtures import strip_unsupported_package_parts


class DocxExtractionTests(SimpleTestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "Megatite S.docx"

    def _create_document(self):
        document = WordDocument()
        document.add_paragraph("معرفی چسب مگاتایت اس")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "دما"
        table.cell(0, 1).text = "زمان"
        table.cell(1, 0).text = "۲۵ درجه"
        table.cell(1, 1).text = "۱۲ ساعت"
        document.add_paragraph("پایان اطلاعات فنی و دستور مصرف محصول")
        document.save(self.path)
        strip_unsupported_package_parts(self.path)

    def test_preserves_paragraph_and_table_order(self):
        self._create_document()

        result = extract_docx(self.path)

        content = result.pages[0].text
        self.assertLess(content.index("معرفی"), content.index("جدول:"))
        self.assertLess(content.index("جدول:"), content.index("پایان"))
        self.assertIn("دما | زمان", content)
        self.assertIn("25 درجه | 12 ساعت", content)
        self.assertEqual(result.block_count, 3)
        self.assertEqual(result.table_count, 1)

    def test_rejects_content_that_would_be_silently_ignored(self):
        self._create_document()
        with ZipFile(self.path, "a") as archive:
            archive.writestr("word/media/technical-facts.txt", "hidden fact")

        with self.assertRaises(DocxExtractionError):
            extract_docx(self.path)

    def test_rejects_hidden_custom_xml_and_thumbnail_parts(self):
        hidden_parts = (
            ("customXml/item1.xml", "<facts>hidden technical fact</facts>"),
            ("docProps/thumbnail.jpeg", b"hidden-thumbnail"),
        )

        for part_name, content in hidden_parts:
            with self.subTest(part=part_name):
                self._create_document()
                with ZipFile(self.path, "a") as archive:
                    archive.writestr(part_name, content)

                with self.assertRaises(DocxExtractionError):
                    extract_docx(self.path)

    def test_rejects_duplicate_archive_entries(self):
        self._create_document()
        with ZipFile(self.path, "a") as archive:
            archive.writestr("word/document.xml", "<duplicate />")

        with self.assertRaises(DocxExtractionError):
            extract_docx(self.path)

    def test_rejects_external_package_relationships(self):
        self._create_document()
        relationship_xml = """<?xml version="1.0" encoding="UTF-8"?>
        <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
          <Relationship Id="hidden" Type="urn:test" Target="https://example.invalid/" TargetMode="External"/>
        </Relationships>
        """
        with ZipFile(self.path, "a") as archive:
            archive.writestr("word/_rels/hidden.rels", relationship_xml)

        with self.assertRaises(DocxExtractionError):
            extract_docx(self.path)

    def test_rejects_malformed_docx(self):
        self.path.write_bytes(b"not-an-office-package")

        with self.assertRaises(DocxExtractionError):
            extract_docx(self.path)
