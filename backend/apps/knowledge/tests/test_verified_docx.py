from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from docx import Document as WordDocument

from apps.knowledge.models import DocumentChunk, DocumentVersion
from apps.knowledge.verified import import_verified_docx

from .docx_fixtures import strip_unsupported_package_parts


class VerifiedDocxTests(TestCase):
    def setUp(self):
        self.media = TemporaryDirectory()
        self.addCleanup(self.media.cleanup)
        self.settings_override = override_settings(MEDIA_ROOT=self.media.name)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.source_directory = TemporaryDirectory()
        self.addCleanup(self.source_directory.cleanup)
        self.path = Path(self.source_directory.name) / "Megatite S.docx"

    def _create_document(self, *, working_time="45"):
        document = WordDocument()
        document.add_paragraph("اطلاعات فنی تاییدشده چسب Megatite S")
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "مشخصه"
        table.cell(0, 1).text = "مقدار"
        table.cell(1, 0).text = "نسبت ترکیب"
        table.cell(1, 1).text = "1 به 1 حجمی"
        table.cell(2, 0).text = "زمان کاربری در دمای 25 درجه"
        table.cell(2, 1).text = f"{working_time} دقیقه"
        document.add_paragraph("این متن و جدول به صورت دستی با دیتاشیت اصلی بررسی شده‌اند.")
        document.save(self.path)
        strip_unsupported_package_parts(self.path)

    @patch("apps.knowledge.ingestion.create_embeddings")
    def test_imports_verified_docx_deduplicates_and_activates_new_version(self, embeddings):
        embeddings.side_effect = lambda texts: [[0.1] * 1024 for _ in texts]
        self._create_document()

        first = import_verified_docx(
            self.path,
            source_key="verified:megatite-s-technical-fa",
            title="Megatite S — اطلاعات فنی تاییدشده",
            embed=True,
        )
        duplicate = import_verified_docx(
            self.path,
            source_key="verified:megatite-s-technical-fa",
            embed=True,
        )

        self.assertTrue(first.created)
        self.assertTrue(duplicate.duplicate)
        self.assertTrue(first.version.content_verified)
        self.assertFalse(first.version.ocr_used)
        self.assertTrue(first.version.is_active)
        self.assertEqual(first.version.status, DocumentVersion.Status.READY)
        self.assertEqual(first.version.original_filename, "Megatite S.docx")
        content = "\n".join(DocumentChunk.objects.values_list("content", flat=True))
        self.assertIn("1 به 1 حجمی", content)
        self.assertIn("45 دقیقه", content)

        self._create_document(working_time="46")
        second = import_verified_docx(
            self.path,
            source_key="verified:megatite-s-technical-fa",
            embed=True,
        )
        first.version.refresh_from_db()
        second.version.refresh_from_db()
        self.assertFalse(first.version.is_active)
        self.assertTrue(second.version.is_active)
        self.assertEqual(second.version.version_number, 2)

    def test_command_requires_explicit_review_confirmation(self):
        self._create_document()

        with self.assertRaises(CommandError):
            call_command("import_verified_docx", str(self.path), stdout=StringIO())

    def test_command_can_safely_dry_run_without_confirmation(self):
        self._create_document()
        output = StringIO()

        call_command(
            "import_verified_docx",
            str(self.path),
            dry_run=True,
            sample_characters=100,
            stdout=output,
        )

        self.assertIn('"tables": 1', output.getvalue())
        self.assertEqual(DocumentVersion.objects.count(), 0)
